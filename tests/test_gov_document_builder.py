from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from scripts.build_gov_documents import build_main_document


ROOT = Path(__file__).resolve().parents[1]


def test_main_document_builder_creates_valid_docx(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    assert output.stat().st_size > 100_000
    with ZipFile(output) as package:
        assert "word/document.xml" in package.namelist()
        document_xml = package.read("word/document.xml").decode("utf-8")

    forbidden_markers = (
        "PREENCHIMENTO OBRIGATÓRIO PELO GRUPO",
        "SEÇÃO AUTORAL — NÃO PREENCHIDA POR IA",
        "[PREENCHER]",
        "[ESCREVER AQUI]",
    )
    assert all(marker not in document_xml for marker in forbidden_markers)


def test_supervision_matrix_uses_landscape_and_keeps_scenarios_together(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    document = Document(output)
    assert any(section.page_width > section.page_height for section in document.sections)

    matrix = next(table for table in document.tables if len(table.columns) == 7)
    assert all(
        cell._tc.get_or_add_tcPr().find(qn("w:noWrap")) is None
        for cell in matrix.rows[0].cells
    )
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for row in matrix.rows[1:]
    )
    assert all(
        paragraph.paragraph_format.first_line_indent == 0
        for row in matrix.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )


def test_final_considerations_start_on_a_new_page(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    document = Document(output)
    heading_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text == "7 CONSIDERAÇÕES FINAIS"
    )
    previous_paragraph_xml = document.paragraphs[heading_index - 1]._p.xml
    assert 'w:type="page"' in previous_paragraph_xml


def test_scenario_table_keeps_score_values_on_one_line(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    document = Document(output)
    scenario_table = next(
        table
        for table in document.tables
        if [cell.text for cell in table.rows[0].cells]
        == [
            "Cenário",
            "Leituras apresentadas",
            "Score",
            "Persistência",
            "Confiança",
            "Resultado governado",
        ]
    )
    score_width = int(
        scenario_table._tbl.tblGrid.gridCol_lst[2].get(qn("w:w"))
    )

    assert score_width >= 800


def test_living_document_preserves_detailed_sprint_1_controls(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    document = Document(output)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )

    for required_control in (
        "Técnico de Operação",
        "photo_file_hash",
        "cv_model_version",
        "fields_low_confidence",
        "Viés de Marca",
        "Model Card",
    ):
        assert required_control in text


def test_living_document_catalogs_sprint_2_mockup_evidence(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Mockup navegável da planta baixa na Sprint 2" in text
    assert "Fonte: Elaborado pelos autores a partir do mockup da Sprint 2 (2025)." in text


def test_document_builder_cli_is_portable_for_main_document(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "build_gov_documents.py"),
            "--main-output",
            str(output),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert output.stat().st_size > 100_000


def test_metric_contract_cli_requires_an_explicit_template(tmp_path):
    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "build_gov_documents.py"),
            "--metric-output",
            str(tmp_path / "metric_contracts.docx"),
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode != 0
    assert "--metric-template" in completed.stderr


def test_page_number_is_hidden_from_cover_and_pretextual_pages(tmp_path):
    output = tmp_path / "challenge_sprint3_gov.docx"

    build_main_document(output)

    document = Document(output)
    assert "PAGE" not in document.sections[0].footer._element.xml
    assert "PAGE" not in document.sections[1].footer._element.xml

    textual_section = document.sections[2]
    assert "PAGE" in textual_section.footer._element.xml
    page_number_type = textual_section._sectPr.find(qn("w:pgNumType"))
    assert page_number_type is not None
    assert page_number_type.get(qn("w:start")) == "4"
    assert all(
        section._sectPr.find(qn("w:pgNumType")) is None
        for section in document.sections[3:]
    )
